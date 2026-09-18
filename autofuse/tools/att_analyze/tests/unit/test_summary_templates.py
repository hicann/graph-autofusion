#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# ----------------------------------------------------------------------------------------------------------
import unittest
import os
import sys
import csv
import io
import tempfile
from contextlib import redirect_stdout
from unittest import mock

sys.path.insert(
    0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "../../src")
)
from summary_templates import LogParser, OperatorSummary
import summary_templates


class TestLogParser(unittest.TestCase):
    def setUp(self):
        self.parser = LogParser()
        # 测试数据在 tests/data/ 目录下
        self.test_dir = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "..", "data"
        )

    def test_extract_operator_names(self):
        with open(
            os.path.join(self.test_dir, "test_concat.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        operators = self.parser.extract_operator_names(log_content)
        self.assertEqual(operators, ["Concat"])

    def test_extract_graph_result(self):
        with open(
            os.path.join(self.test_dir, "test_concat.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        graph, result = self.parser.extract_graph_result(log_content, "Concat")
        self.assertEqual(graph, 0)
        self.assertEqual(result, 0)

    def test_extract_graph_result_missing(self):
        with open(
            os.path.join(self.test_dir, "test_missing_graph.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        graph, result = self.parser.extract_graph_result(log_content, "Subtract")
        self.assertEqual(graph, 0)
        self.assertEqual(result, 0)

    def test_extract_group_case(self):
        with open(
            os.path.join(self.test_dir, "test_concat.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        group, case = self.parser.extract_group_case(log_content, "Concat", 0, 0)
        self.assertEqual(group, 0)
        self.assertEqual(case, 0)

    def test_extract_group_case_missing(self):
        with open(
            os.path.join(self.test_dir, "test_missing_group.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        group, case = self.parser.extract_group_case(log_content, "Divide", 0, 0)
        self.assertEqual(group, -1)
        self.assertEqual(case, 0)

    def test_extract_tiling_values(self):
        with open(
            os.path.join(self.test_dir, "test_concat.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        tiling_values = self.parser.extract_tiling_values(
            log_content, "Concat", 0, 0, 0, 0
        )
        self.assertEqual(tiling_values["s0t_size"], 256)
        self.assertEqual(tiling_values["s1Ts0Tb_size"], 4096)
        self.assertEqual(tiling_values["s1t_size"], 1)
        self.assertEqual(tiling_values["ub_size"], 1024)
        self.assertEqual(tiling_values["block_dim"], 1)
        self.assertEqual(tiling_values["q0_size"], 1024)
        self.assertEqual(tiling_values["AIV_MTE2"], 387906.099166)
        self.assertEqual(tiling_values["AIV_MTE3"], 355058.291519)

    def test_extract_performance_metrics(self):
        with open(
            os.path.join(self.test_dir, "test_concat.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        metrics = self.parser.extract_performance_metrics(
            log_content, "Concat", 0, 0, 0, 0
        )
        self.assertEqual(metrics["objective_value"], 388206.099166)
        self.assertEqual(metrics["aiv_mte2"], 387906.099166)
        self.assertEqual(metrics["aiv_mte3"], 355058.291519)

    def test_extract_result_performance(self):
        with open(
            os.path.join(self.test_dir, "test_complete.log"), "r", encoding="utf-8"
        ) as f:
            log_content = f.read()

        result_perf = self.parser.extract_result_performance(log_content, "Add", 0, 0)
        self.assertEqual(result_perf, 194103.049583)

    def test_extract_result_performance_missing(self):
        with open(
            os.path.join(self.test_dir, "test_missing_result_perf.log"),
            "r",
            encoding="utf-8",
        ) as f:
            log_content = f.read()

        result_perf = self.parser.extract_result_performance(
            log_content, "MatMul", 0, 0
        )
        self.assertIsNone(result_perf)

    def test_parse_log_file_concat(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_concat.log")
        )
        self.assertEqual(len(summaries), 1)

        summary = summaries[0]
        self.assertEqual(summary.operator_name, "Concat")
        self.assertEqual(summary.graph, 0)
        self.assertEqual(summary.result, 0)
        self.assertEqual(summary.group, 0)
        self.assertEqual(summary.case, 0)
        self.assertEqual(summary.aiv_mte2, 387906.099166)
        self.assertEqual(summary.aiv_mte3, 355058.291519)
        self.assertEqual(summary.objective_value, 388206.099166)
        self.assertIsNone(summary.result_performance)
        self.assertEqual(summary.tiling_values["s0t_size"], 256)
        self.assertEqual(summary.tiling_values["s1Ts0Tb_size"], 4096)
        self.assertEqual(summary.tiling_values["s1t_size"], 1)
        self.assertEqual(summary.tiling_values["ub_size"], 1024)
        self.assertEqual(summary.tiling_values["block_dim"], 1)
        self.assertEqual(summary.tiling_values["q0_size"], 1024)

    def test_parse_log_file_complete(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_complete.log")
        )
        self.assertEqual(len(summaries), 2)

        summaries_dict = {s.operator_name: s for s in summaries}

        add_summary = summaries_dict["Add"]
        self.assertEqual(add_summary.operator_name, "Add")
        self.assertEqual(add_summary.graph, 0)
        self.assertEqual(add_summary.result, 0)
        self.assertEqual(add_summary.group, 0)
        self.assertEqual(add_summary.case, 0)
        self.assertEqual(add_summary.aiv_mte2, 193953.049583)
        self.assertEqual(add_summary.aiv_mte3, 177529.145759)
        self.assertEqual(add_summary.objective_value, 194103.049583)
        self.assertEqual(add_summary.result_performance, 194103.049583)

        mul_summary = summaries_dict["Mul"]
        self.assertEqual(mul_summary.operator_name, "Mul")
        self.assertEqual(mul_summary.graph, 1)
        self.assertEqual(mul_summary.result, 1)
        self.assertEqual(mul_summary.group, 0)
        self.assertEqual(mul_summary.case, 0)
        self.assertEqual(mul_summary.aiv_mte2, 96976.524791)
        self.assertEqual(mul_summary.aiv_mte3, 88764.572879)
        self.assertEqual(mul_summary.objective_value, 97051.524791)
        self.assertEqual(mul_summary.result_performance, 97051.524791)

    def test_parse_log_file_missing_graph(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_missing_graph.log")
        )
        self.assertEqual(len(summaries), 1)

        summary = summaries[0]
        self.assertEqual(summary.operator_name, "Subtract")
        self.assertEqual(summary.graph, 0)
        self.assertEqual(summary.result, 0)
        self.assertEqual(summary.group, 0)
        self.assertEqual(summary.case, 0)

    def test_parse_log_file_missing_group(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_missing_group.log")
        )
        self.assertEqual(len(summaries), 1)

        summary = summaries[0]
        self.assertEqual(summary.operator_name, "Divide")
        self.assertEqual(summary.graph, 0)
        self.assertEqual(summary.result, 0)
        self.assertEqual(summary.group, -1)
        self.assertEqual(summary.case, 0)

    def test_parse_log_file_missing_result_perf(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_missing_result_perf.log")
        )
        self.assertEqual(len(summaries), 1)

        summary = summaries[0]
        self.assertEqual(summary.operator_name, "MatMul")
        self.assertEqual(summary.graph, 0)
        self.assertEqual(summary.result, 0)
        self.assertEqual(summary.group, 0)
        self.assertEqual(summary.case, 0)
        self.assertIsNone(summary.result_performance)

    def test_multiple_tiling_values(self):
        summaries = self.parser.parse_log_file(
            os.path.join(self.test_dir, "test_missing_result_perf.log")
        )
        self.assertEqual(len(summaries), 1)

        summary = summaries[0]
        self.assertIn("s0t_size", summary.tiling_values)
        self.assertIn("s1Ts0Tb_size", summary.tiling_values)
        self.assertIn("s1t_size", summary.tiling_values)
        self.assertIn("s2t_size", summary.tiling_values)
        self.assertIn("ub_size", summary.tiling_values)
        self.assertIn("block_dim", summary.tiling_values)
        self.assertIn("q0_size", summary.tiling_values)
        self.assertIn("q1_size", summary.tiling_values)


class TestSummaryMainAllFlag(unittest.TestCase):
    @mock.patch("summary_templates.print_summary_table")
    @mock.patch("summary_templates.find_log_files")
    @mock.patch("summary_templates.LogParser")
    def test_summary_main_all_flag_passes_all_results_mode(
        self, mock_parser_cls, mock_find_log_files, _mock_print
    ):
        mock_find_log_files.return_value = ["/tmp/fake.log"]
        mock_parser = mock_parser_cls.return_value
        mock_parser.parse_log_file.return_value = [
            OperatorSummary(operator_name="Fake")
        ]

        with (
            mock.patch("summary_templates.os.path.exists", return_value=True),
            mock.patch("summary_templates.os.path.isdir", return_value=True),
            mock.patch.object(sys, "argv", ["summary_templates", "/tmp/logs", "--all"]),
        ):
            summary_templates.main()

        mock_parser.parse_log_file.assert_called_once_with(
            "/tmp/fake.log", summary_mode="all_results_all_groups"
        )


class TestFinalTilingExport(unittest.TestCase):
    def setUp(self):
        self.test_dir = os.path.join(os.path.dirname(__file__), "..", "data")

    def test_final_records_are_preferred_and_append_columns(self):
        output = os.path.join(self.test_dir, "_final_summary.csv")
        try:
            path = os.path.join(self.test_dir, "test_final_tiling_multigroup.log")
            with open(path, encoding="utf-8") as stream:
                content = stream.read()
            records = LogParser().extract_final_tiling_records(content, path)
            summary_templates.export_to_csv([], output, final_records=records)
            with open(output, newline="", encoding="utf-8-sig") as stream:
                rows = list(csv.reader(stream))
            self.assertEqual(len(rows), 4)
            self.assertEqual(
                rows[0][-6:],
                [
                    "Final Source",
                    "Tiling Key",
                    "Score",
                    "Pipe Estimate",
                    "Tiling Repr",
                    "Final Parse Status",
                ],
            )
            self.assertEqual(
                rows[1][-6:],
                [
                    "r",
                    "2",
                    "",
                    '{"AIV_MTE2":128.0,"AIV_MTE3":null,"AIV_VEC":null}',
                    '{"tile_m":64}',
                    "ok",
                ],
            )
            self.assertEqual(
                {(row[0], row[3]) for row in rows[1:]},
                {("Fusion_0", "0"), ("Fusion_0", "1"), ("Fusion_1", "0")},
            )
        finally:
            if os.path.exists(output):
                os.unlink(output)

    def test_final_records_are_exported_to_excel_without_legacy_summary(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        path = os.path.join(self.test_dir, "_final_summary.xlsx")
        try:
            with open(
                os.path.join(self.test_dir, "test_final_tiling_single.log"),
                encoding="utf-8",
            ) as stream:
                records = LogParser().extract_final_tiling_records(stream.read())
            summary_templates.export_to_excel([], path, final_records=records)
            workbook = openpyxl.load_workbook(path, read_only=True)
            headers = [cell.value for cell in next(workbook.active.iter_rows())]
            self.assertIn("Tiling Repr", headers)
            self.assertEqual(workbook.active.cell(row=2, column=1).value, "Add_0")
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_mixed_final_and_legacy_rows_retain_historical_tiling_values(self):
        content = (
            "[Fusion] [PROF]The value of foo is 7 in graph0_result0_g0_2.\n"
            "[Fusion] [PROF]The value of ub_size is 8 in graph0_result0_g0_2.\n"
            "[Fusion] [PROF]The value of block_dim is 9 in graph0_result0_g0_2.\n"
            "[Fusion] [PROF]Among the templates, tiling case 2 of graph0_result0_g0 is the best choice.\n"
            "[Fusion] [PROF]Among all schedule results, graph0_result0 is the best choice.\n"
            '[ATT][FINAL_TILING] schema=1 source=r operator=Fusion graph=0 result=0 group=0 case_id=2 tiling_key=2 pipe_estimates="{}" tiling_repr="{\\"tile_m\\":64}"\n'
        )
        with tempfile.TemporaryDirectory() as directory:
            log_path = os.path.join(directory, "mixed.log")
            output = os.path.join(directory, "summary.csv")
            with open(log_path, "w", encoding="utf-8") as stream:
                stream.write(content)
            parser = LogParser()
            summaries = parser.parse_log_file(log_path)
            records = parser.extract_final_tiling_records(content, log_path)
            summary_templates.export_to_csv(summaries, output, final_records=records)
            with open(output, newline="", encoding="utf-8-sig") as stream:
                rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["foo"], "7.0")
        self.assertEqual(rows[0]["ub_size"], "8.0")
        self.assertEqual(rows[0]["block_dim"], "9.0")

    def test_summary_main_console_prints_final_records(self):
        log_path = os.path.join(self.test_dir, "test_final_tiling_multigroup.log")
        output = io.StringIO()
        with (
            mock.patch.object(sys, "argv", ["summary_templates", log_path]),
            redirect_stdout(output),
        ):
            summary_templates.main()

        rendered = output.getvalue()
        self.assertIn("Final Source", rendered)
        self.assertIn("Fusion_0", rendered)
        self.assertIn('tile_m":64', rendered)
        self.assertIn('AIV_MTE2":128.0', rendered)

    @mock.patch("summary_templates.print_summary_table")
    @mock.patch("summary_templates.LogParser")
    def test_summary_main_default_passes_best_result_all_groups_mode(
        self, mock_parser_cls, _mock_print
    ):
        mock_parser = mock_parser_cls.return_value
        mock_parser.parse_log_file.return_value = [
            OperatorSummary(operator_name="Fake")
        ]

        with mock.patch.object(
            sys,
            "argv",
            [
                "summary_templates",
                os.path.join(
                    os.path.dirname(__file__), "..", "data", "test_concat.log"
                ),
            ],
        ):
            summary_templates.main()

        mock_parser.parse_log_file.assert_called_once_with(
            os.path.join(os.path.dirname(__file__), "..", "data", "test_concat.log"),
            summary_mode="best_result_all_groups",
        )


if __name__ == "__main__":
    unittest.main()
